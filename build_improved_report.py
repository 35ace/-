"""Build a revision using the original report's model and document styles."""
import csv
import json
import math
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from report_styles import (add_body, add_caption, add_heading, add_table,
                           remove_paragraph_borders, remove_style_borders, set_font)

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output" / "improved_lqr"
SOURCE = ROOT / "飞行器控制算法设计项目报告书初稿.docx"
DEST = ROOT / "飞行器控制算法设计项目报告书改进版.docx"


def rows(name):
    with (OUT / name).open(encoding="utf-8-sig", newline="") as stream:
        return list(csv.DictReader(stream))


def number(value, digits=2):
    value = float(value)
    return "未在时窗内调定" if not math.isfinite(value) else f"{value:.{digits}f}"


def equation(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(5)
    p.paragraph_format.keep_together = True
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_font(run, size=10.5)
    run.font.name = "Cambria"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Cambria")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Cambria")
    return p


def figure(doc, filename, caption, width=15.0):
    path = OUT / filename
    assert path.exists(), path
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.keep_with_next = True
    p.add_run().add_picture(str(path), width=Cm(width))
    add_caption(doc, caption)


def page(doc, heading):
    doc.add_page_break()
    add_heading(doc, heading)


def comparison_rows(metrics, scenario, methods):
    return [[m, number(metrics[(scenario, m)]["MaxPitch_deg"]),
             number(metrics[(scenario, m)]["SettlingTime_s"]),
             number(metrics[(scenario, m)]["TargetIAE_m_s"]),
             number(metrics[(scenario, m)]["MaxElevator_deg"])] for m in methods]


def build():
    comparison = rows("comparison_metrics.csv")
    metrics = {(r["Scenario"], r["Method"]): r for r in comparison}
    original = rows("original_reproduction.csv")[0]
    gain = rows("lqr_gain.csv")[0]
    analysis = json.loads((OUT / "analysis_summary.json").read_text(encoding="utf-8"))
    config = json.loads((OUT / "configuration.json").read_text(encoding="utf-8"))
    integration = rows("integration_check.csv")[0]
    full_name = "SAFE_RG_DOB_AW"
    nominal = metrics[("nominal_gust", full_name)]
    baseline = metrics[("nominal_gust", "LQI")]
    full_mc = analysis["methods"][full_name]
    baseline_mc = analysis["methods"]["LQI"]

    # Preserve the original styles and page settings, and save a distinct revision.
    doc = Document(SOURCE)
    for child in list(doc._element.body):
        if child.tag != qn("w:sectPr"):
            doc._element.body.remove(child)
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    for style_name in ["Title", "Subtitle", "Heading 1", "Heading 2"]:
        doc.styles[style_name].font.color.rgb = RGBColor(0, 0, 0)
        remove_style_borders(doc.styles[style_name])
    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_paragraph_borders(title)
    set_font(title.add_run("飞行器高度姿态控制算法\n设计与仿真项目报告书"), size=20, bold=True)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(p.add_run("中南大学第七届先进飞行器虚拟仿真大赛 选题五"), size=12)
    add_table(doc, ["项目条目", "内容"], [
        ["项目名称", "基于扰动观测补偿与俯仰包线约束的积分 LQR 高度控制"],
        ["控制对象", "原报告的四状态纵向模型及升降舵通道"],
        ["仿真平台", "MATLAB R2025a 及 Control System Toolbox"],
        ["提交信息", "参赛队名称、单位及队员信息由参赛队补充"],
    ], [3.0, 13.0])
    add_heading(doc, "摘要")
    add_body(doc, "本项目沿用纵向小扰动模型及积分 LQR 反馈增益，针对高度捕获时俯仰角过大、阵风扰动及执行器饱和问题，设计扰动观测补偿、有限时域参考治理、俯仰包线安全滤波和积分抗饱和的组合控制结构。扰动从垂向加速度通道进入，与升降舵输入不匹配，因此通过稳态输出调节方程求取补偿系数；俯仰包线控制作用于舵机速度，抗饱和反馈同步考虑舵机限幅、限速及安全修正。")
    add_body(doc, f"保留原算法作为基线，在相同舵机模型下进行九组方法、六种工况及30组随机参数对比。标称阵风工况俯仰峰值从{number(baseline['MaxPitch_deg'])}°降至{number(nominal['MaxPitch_deg'])}°，但高度捕获时间从{number(baseline['SettlingTime_s'])} s增加至{number(nominal['SettlingTime_s'])} s。最终随机验证中完整方案最差俯仰峰值为{number(full_mc['MaxPitch_deg']['max'])}°，{full_mc['pitch_exceedance_count']}/30组超过10°。结果支持降低测试范围内的姿态越界风险，不能据此证明真实飞行器全包线安全。")

    page(doc, "一 设计目标与原方案复现")
    add_body(doc, "原报告采用高度偏差 h、垂向速度 v、俯仰角 θ、俯仰角速度 q 构成纵向模型，以升降舵偏角 u 为输入。高度目标为100 m，经7 s一阶滤波，14至20 s加入−0.8 m/s²阵风，舵面限制为±25°。本项目保留该控制对象、A与B矩阵、Q与R权重，并将原积分 LQR 称为 LQI。")
    add_table(doc, ["原45秒工况指标", "独立复现值"], [
        ["45 s末端高度误差", number(original["FinalError_m"], 3) + " m"],
        ["2 m误差带调节时间", number(original["SettlingTime_s"], 3) + " s"],
        ["最大俯仰角", number(original["MaxPitch_deg"], 3) + "°"],
        ["最大升降舵偏角", number(original["MaxElevator_deg"], 3) + "°"],
        ["目标高度绝对误差积分", number(original["TargetIAE_m_s"], 3) + " m·s"],
        ["非负峰值超调", number(original["Overshoot_m"], 3) + " m"],
    ], [8.0, 8.0])
    add_body(doc, "原方案的29.10°俯仰峰值较大，使小扰动模型外推的可信度下降。原报告将 max(h)−100 写为负超调，修订为 max(0,max(h)−100)；45 s末端误差改称有限时刻误差，避免把一次仿真末值当成稳态证明。单一无显著饱和工况也不足以评价抗饱和设计。")
    add_body(doc, "本项目设置10°为算法验证中的俯仰评价阈值。它用于检验姿态峰值是否得到约束，没有真实机型包线依据。改进方案将内部安全滤波阈值设为6°，保留4°裕量，面对模型摄动与测量噪声进行数值验证。")
    add_body(doc, "设计目标是以可解释的响应速度代价降低俯仰峰值，并对持续阵风、气动参数偏差、执行器限速及高度指令反向进行验证。所有算法采用同一模型、同一扰动和同一噪声实现；除专门的原脚本复现外，对比组均包含相同舵机动态。")

    page(doc, "二 控制对象与执行器建模")
    add_body(doc, "沿用原报告的线性模型 ẋ = Ax + Bu + Ed，x = [h,v,θ,q]ᵀ，E = [0,1,0,0]ᵀ。d为等效垂向加速度扰动。状态变量和输入均为配平点附近偏差；模型参数尚未对应具体飞行器气动导数。")
    add_table(doc, ["A矩阵行", "h", "v", "θ", "q", "B", "E"], [
        ["ḣ", "0", "1", "0", "0", "0", "0"],
        ["v̇", "0", "−0.42", "8.50", "0.35", "0.85", "1"],
        ["θ̇", "0", "0", "0", "1", "0", "0"],
        ["q̇", "0", "−0.030", "−1.85", "−1.20", "2.40", "0"],
    ], [2.6, 1.7, 2.4, 2.4, 2.4, 2.3, 2.2])
    add_body(doc, "高度单位为m，垂向速度为m/s，角度为rad，角速度为rad/s；报告图表转换为度。垂向阵风施加于E通道，与B通道不共线，不能通过简单的输入扰动抵消公式精确消除。")
    add_heading(doc, "舵机动态与可实现输入", 2)
    equation(doc, "u̇ = clip((clip(u_c, −u_max, u_max) − u)/τ, −ρ, ρ)")
    add_body(doc, "u_c为控制器原始命令，u为实际舵角。标称时间常数τ=0.15 s，舵角限幅±25°，速率限制±60°/s；压力工况采用τ=0.30 s、±15°和±12°/s。安全滤波会进一步修正实际舵机速度。")
    add_heading(doc, "测量与仿真范围", 2)
    add_body(doc, "有噪声工况对[h,v,θ,q]分别加入标准差[0.03 m,0.025 m/s,0.05°,0.08°/s]的独立零均值高斯噪声。每0.02 s生成并保持一个样本；该设置是算法压力测试，不是经标定的传感器模型。控制律、扰动观测器和参考治理只读取测量值，真实状态用于动力学和性能评价。")
    add_body(doc, "该四状态模型没有空速、推力、迎角或气动非线性，因此不能直接评价失速、能量管理及真实飞行器100 m爬升可行性。模拟姿态越界的基线仅用于算法对照，超出小扰动范围时不解释为真实飞行性能。")

    page(doc, "三 文献依据与算法结构")
    add_body(doc, "参考文献用于方法基础，各模块按本项目控制对象重新推导和实现。四旋翼自抗扰论文[3]提供观测与补偿的启发；本项目采用固定翼纵向通道，没有沿用四旋翼动力学、控制参数或其仿真结论。")
    add_table(doc, ["方法依据", "借鉴内容", "本项目实现"], [
        ["Chen等[4]", "扰动观测控制", "垂向速度降阶观测器及非匹配扰动稳态补偿"],
        ["Garone等[5]", "参考与指令治理", "12 s闭环预测及一维参考增量约束"],
        ["Kothare等[6]", "积分抗饱和", "对限幅、限速及安全修正后的可实现输入回算"],
        ["Xiao与Belta[7]", "高阶控制障碍函数", "包含舵机动态的三阶俯仰包线滤波"],
    ], [3.2, 4.0, 8.8])
    add_heading(doc, "闭环信号关系", 2)
    add_body(doc, "目标高度首先经过参考治理生成r；测量状态与r进入积分 LQR，叠加扰动观测补偿得到u_c。舵角限幅、舵速限制及俯仰安全滤波生成u̇，再积分得到实际舵角u作用于飞行器。实际可实现输入同时返回积分抗饱和环节，观测器也使用实际u，避免把执行器饱和误判成外部扰动。")
    add_table(doc, ["代码方法名", "参考治理", "扰动补偿", "抗饱和", "安全滤波"], [
        ["LQI", "无", "无", "无", "无"],
        ["LQI_AW", "无", "无", "有", "无"],
        ["RG_AW", "有", "无", "有", "无"],
        ["RG_DOB", "有", "有", "无", "无"],
        ["RG_DOB_AW", "有", "有", "有", "无"],
        ["DOB_AW_SAFE", "无", "有", "有", "有"],
        ["SAFE_RG_DOB_AW", "有", "有", "有", "有"],
        ["RG_AW_SAFE", "有", "无", "有", "有"],
        ["RG_DOB_SAFE", "有", "有", "无", "有"],
    ], [5.8, 2.55, 2.55, 2.55, 2.55])
    add_body(doc, "LQI及没有参考治理的算法保留原7 s一阶指令滤波。全部方法保留同一组LQR增益；这样可辨别新增模块的作用。组合结构尚未证明为文献首创，比赛贡献定位为针对原模型的设计改进与验证。")

    page(doc, "四 积分 LQR 与扰动补偿推导")
    add_heading(doc, "积分增广控制", 2)
    equation(doc, "ξ̇ = r − h     x_a = [xᵀ, ξ]ᵀ     u_c = −K_x y − K_i ξ + F_d d̂")
    add_body(doc, "y为测量状态，K=[K_x,K_i]通过原权重Q=diag(0.90,0.30,45,6,0.16)与R=1.15求解。积分状态用于消除持续偏差，补偿项仅在DOB方法中启用。LQR性能指标针对零参考增广模型，不能将它直接解释为全部非零参考和饱和工况的全局最优解。")
    equation(doc, "K = [" + ", ".join(number(gain[k], 6) for k in
            ["K_h", "K_hdot", "K_theta", "K_q", "K_integral"]) + "]")
    add_heading(doc, "避免测量微分的降阶扰动观测器", 2)
    equation(doc, "d̂ = z + ℓ y_v     ż = −ℓ(z + ℓ y_v + A₂ y + B₂ u)")
    add_body(doc, "A₂为A的第二行，B₂=0.85，ℓ=1.5 s⁻¹。标称且无噪声时，d̂̇=ℓ(d−d̂)，常值扰动估计误差按一阶指数衰减，时间常数约0.667 s。扰动突变期间必然存在估计滞后；提高带宽也会放大噪声。模型摄动时观测器估计的是包含第二行模型误差的等效扰动，不能保证识别其他通道的全部不确定性。")
    add_heading(doc, "非匹配扰动的稳态调节补偿", 2)
    equation(doc, "[A  B; C_h  0] [X_d; U_d] = [−E; 0]     F_d = K_x X_d + U_d")
    equilibrium = config["DisturbanceEquilibrium"]
    equation(doc, f"X_d = [0, 0, {number(equilibrium[2], 6)}, 0]ᵀ     U_d = {number(equilibrium[4], 6)}")
    equation(doc, f"F_d = {number(config['DOBFeedforward'], 6)}")
    add_body(doc, "常值d下，x=X_dd、u=U_dd满足高度输出为零的扰动平衡。令u_c=−K_xx−K_iξ+F_dd̂，当d̂=d时，可在ξ不承担补偿的情况下生成相应平衡舵角。这只保证标称稳态调节关系，不能宣称瞬态精确消除阵风；积分环节仍用于补偿残差。")

    page(doc, "五 参考治理与积分抗饱和")
    add_heading(doc, "有限时域参考治理", 2)
    equation(doc, "r_k = r_(k−1) + α(r_target − r_(k−1))     0 ≤ α ≤ 1")
    add_body(doc, "每0.2 s更新一次r。对含舵机的一阶闭环模型做矩阵指数离散化，预测步长0.1 s，时域12 s。预测期间保持候选参考和当前扰动估计不变，约束预测俯仰±8°、实际舵角与命令幅值及舵速。使用同一组LQR增益，控制器不读取未来真实阵风。")
    equation(doc, "ẇ = A_c w + G_r r + G_d d̂     w = [xᵀ, ξ, u]ᵀ")
    equation(doc, "y_pred = P w + G_r,pred r_(k−1) + G_d,pred d̂ + g α")
    add_body(doc, "各预测约束对α均为线性不等式，求交集后选择最大可行α，使参考尽可能向目标推进；不需要在线求解一般二次规划。若交集为空，则在21个候选α中选择最大归一化约束违反量最小者，并记录不可行次数。这是显式回退规则，没有鲁棒不变集或递归可行保证。")
    add_body(doc, "实际安全滤波和抗饱和会改变预测采用的线性闭环，参数摄动也会引起预测失配。因此参考治理作为指令管理层，实际姿态风险由安全滤波及数值验证进一步评估。报告保留不可行计数，而不是将回退行为记为成功满足所有预测约束。")
    add_heading(doc, "考虑实际舵机与安全修正的回算", 2)
    equation(doc, "u_eq = u + τu̇_final     ξ̇ = r − y_h + k_aw(u_eq − u_c)/(−K_i)")
    add_body(doc, "k_aw=2 s⁻¹。u_eq是当前实际舵角与最终可实现舵速反推的等效输入；u̇_final已经包含限幅、限速及安全滤波修正。K_i<0，因此除以−K_i后回算方向与积分驱动舵角一致。当执行器线性且未受安全修正时u_eq=u_c，抗饱和项为零，原积分控制得到保留。")
    add_body(doc, "这种连接避免积分器在安全层限制动作时持续堆积，也避免将正常舵机一阶滞后全部当成饱和。其作用需要在实际限幅和限速工况下检验，标称不饱和工况中LQI与LQI_AW结果相同是合理现象。")

    page(doc, "六 俯仰包线安全滤波")
    add_body(doc, "俯仰角θ对舵机速度ν=u̇的相对阶为三。设名义俯仰角加速度a_θ=A₄y+B₄u，名义俯仰角加加速度漂移j₀=A₄(Ay+Bu+Ed̂)，其中A₄为第四行、B₄=2.40。安全滤波不直接截断θ，而是提前限制舵机速度。")
    equation(doc, "b_+ = θ_s − θ     b_− = θ_s + θ     θ_s = 6°")
    equation(doc, "(D + λ)³ b_± ≥ 0     λ = 2 s⁻¹")
    add_body(doc, "展开三阶指数障碍条件，得到舵速上下界。定义c=−j₀−3λa_θ−3λ²q，则：")
    equation(doc, "ν_low = [c − λ³(θ_s + θ)]/B₄")
    equation(doc, "ν_high = [c + λ³(θ_s − θ)]/B₄")
    add_body(doc, "将该区间与舵速限制[−ρ,ρ]及一步舵角限幅区间求交，投影原舵机速度ν₀到可行区间。对单输入，这是最小化(ν−ν₀)²的区间投影，尽量保留原控制动作。若区间为空，则采用明确的最小偏离回退并记录安全不可行时长。")
    equation(doc, "ν = arg min (ν − ν₀)²     subject to ν ∈ 安全与执行器可行区间")
    add_heading(doc, "适用条件与数值裕量", 2)
    add_body(doc, "障碍条件依赖准确模型、可行输入及高阶初始安全条件。当前从零偏差状态初始化，满足标称初始条件；不能直接推出任意初始状态、任意气动参数及任意阵风下安全。测量噪声、离散采样和估计滞后会使θ略超内部6°阈值，所以以10°评价阈值进行实测验证。")
    add_body(doc, "内部8°阈值在开发参数集上曾出现越过10°，据此收紧为6°，随后使用不同随机参数种子进行最终验证。最终30组样本没有越过10°仍不构成对整个不确定集合的形式证明；需要进一步求取鲁棒不变集或加入可证的模型误差上界。")

    page(doc, "七 仿真工况与评价规则")
    add_body(doc, "改进模型仿真时长120 s，控制计算与噪声保持间隔0.02 s，采用四阶Runge–Kutta积分。原始45 s工况单独使用ode45及原容差复现；不得把45 s原始末值与120 s新末值直接作为改进百分比。")
    add_table(doc, ["工况", "具体设置"], [
        ["标称阵风", "14至20 s，d=−0.8 m/s²；标称舵机；无测量噪声"],
        ["强与持续阵风", "14至20 s，d=−1.2；45至65 s，d=−0.45 m/s²"],
        ["参数与噪声", "第二行导数比例[1.20,0.85,1.15]，第四行[0.85,1.15,0.80]，B降20%；加入测量噪声"],
        ["舵机压力", "τ=0.30 s，幅值±15°、速率±12°/s；14至20 s，d=−1.2"],
        ["指令反向", "舵机压力设置；55 s将高度目标由100 m改为40 m"],
        ["定高抗扰", "高度偏差目标0 m；14至20 s，d=−0.8；45至65 s，d=−0.45 m/s²"],
    ], [3.5, 12.5])
    add_heading(doc, "主要性能指标", 2)
    add_table(doc, ["指标", "定义及比较边界"], [
        ["目标IAE", "∫|r_target−h|dt，所有算法使用相同目标；体现姿态约束付出的速度代价"],
        ["参考RMSE", "sqrt(∫(r−h)²dt/T)，不同算法内部r不同，不能单独排名"],
        ["调节时间", "此后始终进入目标±2 m误差带的最早时刻；反向工况从55 s起算；未调定记NaN"],
        ["姿态及执行器", "俯仰峰值、超10°时长、舵角与舵速峰值、原始命令饱和时长及舵角总变差"],
        ["不可行记录", "参考治理不可行更新数及安全滤波不可行时长；原始命令饱和不等同实际舵角越限"],
    ], [3.5, 12.5])
    add_body(doc, "随机试验对A中第二、第四行的六个非零导数及B的比例在[0.8,1.2]均匀抽样；阵风强度比例在[0.75,1.5]抽样，45至65 s叠加−0.30 m/s²，加入测量噪声。最终参数种子为20262009加组号，噪声种子为2000加组号；同组参数与噪声在算法间严格配对。")

    page(doc, "八 标称阵风与速度代价")
    figure(doc, "nominal_gust_comparison.png", "图1 标称阵风下高度 俯仰与舵角对比", 15.0)
    add_table(doc, ["方法", "俯仰峰值 °", "调节时间 s", "目标IAE m·s", "舵角峰值 °"],
              comparison_rows(metrics, "nominal_gust", ["LQI", "DOB_AW_SAFE", full_name]),
              [5.4, 2.65, 2.65, 2.65, 2.65])
    add_body(doc, f"相同舵机模型下，完整方案俯仰峰值降低{100*(1-float(nominal['MaxPitch_deg'])/float(baseline['MaxPitch_deg'])):.1f}%，调节时间增加{float(nominal['SettlingTime_s'])-float(baseline['SettlingTime_s']):.2f} s。目标IAE也增加，说明本方案获得姿态约束收益的同时牺牲了高度捕获速度。")
    add_body(doc, "图中虚线为共同目标高度，点线为完整方案治理后的参考。蓝色RG_AW仍有姿态振荡，说明单独参考预测不足以约束阵风下的真实姿态。绿色与青色均加入安全滤波，峰值较小；参考治理是否值得保留，需要结合反向机动和参数扰动结果判断。")

    page(doc, "九 强阵风与执行器压力测试")
    figure(doc, "strong_persistent_gust_comparison.png", "图2 强阵风与持续阵风下的闭环响应", 14.4)
    pressure_lqi = metrics[("actuator_stress", "LQI")]
    pressure_aw = metrics[("actuator_stress", "LQI_AW")]
    pressure_full = metrics[("actuator_stress", full_name)]
    add_table(doc, ["舵机压力方法", "俯仰 °", "末端误差 m", "调节时间 s", "目标IAE m·s"], [
        [m, number(metrics[("actuator_stress",m)]["MaxPitch_deg"]),
         number(metrics[("actuator_stress",m)]["FinalError_m"],3),
         number(metrics[("actuator_stress",m)]["SettlingTime_s"]),
         number(metrics[("actuator_stress",m)]["TargetIAE_m_s"])]
        for m in ["LQI", "LQI_AW", full_name]], [5.4, 2.2, 2.9, 2.8, 2.7])
    add_body(doc, "强阵风曲线显示安全滤波抑制阵风补偿过程中的俯仰峰值，持续扰动需改变俯仰及舵角平衡。该工况100 m捕获尚在进行，不能将全部高度误差都解释为抗扰性能；完整方案的慢参考本身也是误差来源。")
    add_body(doc, f"舵机压力工况LQI的120 s末端误差为{number(pressure_lqi['FinalError_m'],3)} m；加入抗饱和后为{number(pressure_aw['FinalError_m'],3)} m，但俯仰仍为{number(pressure_aw['MaxPitch_deg'])}°。完整方案俯仰降至{number(pressure_full['MaxPitch_deg'])}°，调节时间为{number(pressure_full['SettlingTime_s'])} s，显示抗饱和与姿态约束解决不同问题。")

    page(doc, "十 指令反向与模块消融")
    figure(doc, "command_reversal_comparison.png", "图3 55秒高度目标由100米改为40米", 14.7)
    add_table(doc, ["方法", "俯仰峰值 °", "反向调节 s", "目标IAE m·s", "舵角峰值 °"],
              comparison_rows(metrics, "command_reversal", ["LQI_AW", "DOB_AW_SAFE", full_name]),
              [5.4, 2.65, 2.65, 2.65, 2.65])
    reversal = metrics[("command_reversal", full_name)]
    no_rg = metrics[("command_reversal", "DOB_AW_SAFE")]
    add_body(doc, f"完整方案反向后{number(reversal['SettlingTime_s'])} s进入并保持±2 m误差带，去掉参考治理后为{number(no_rg['SettlingTime_s'])} s；完整方案的舵角总变差为{number(reversal['ElevatorTotalVariation_deg'])}°，去掉参考治理后为{number(no_rg['ElevatorTotalVariation_deg'])}°。在该工况参考治理有助于管理反向命令，但不代表每个工况都更快。")
    add_body(doc, "消融结果中，标称LQI与LQI_AW相同，因为未发生显著限幅；舵机压力下抗饱和改善恢复。RG_DOB没有抗饱和时在压力与反向工况出现较大误差，说明观测补偿不能代替积分管理。去掉安全滤波的RG_DOB_AW仍越过10°，说明单靠观测与预测不能满足所设评价阈值。完整数据保留所有九组，避免仅展示最有利对照。")

    page(doc, "十一 定高抗扰与逐项消融")
    figure(doc, "disturbance_observer.png", "图4 定高工况扰动估计及高度偏差", 14.5)
    hold = metrics[("altitude_hold_gust", full_name)]
    no_dob = metrics[("altitude_hold_gust", "RG_AW_SAFE")]
    add_table(doc, ["定高方法", "俯仰峰值 °", "最大偏差 m", "目标IAE m·s", "持续阵风RMSE m"], [
        [m,number(metrics[("altitude_hold_gust",m)]["MaxPitch_deg"]),
         number(metrics[("altitude_hold_gust",m)]["MaxAbsoluteTargetError_m"],3),
         number(metrics[("altitude_hold_gust",m)]["TargetIAE_m_s"],3),
         number(metrics[("altitude_hold_gust",m)]["GustWindowTargetRMSE_m"],3)]
        for m in ["LQI", "RG_DOB_AW", "RG_AW_SAFE", full_name]],
        [5.2,2.5,2.5,2.6,3.2])
    add_body(doc, f"在相同安全与抗饱和结构下，补偿将45至65 s持续阵风窗口的高度误差RMSE从{number(no_dob['GustWindowTargetRMSE_m'],3)} m降至{number(hold['GustWindowTargetRMSE_m'],3)} m。安全滤波压制了抵抗第一段阵风所需的约7°俯仰，因此完整方案全时域高度IAE仍高于无安全滤波的LQI；姿态与抗扰能力之间存在物理约束取舍。")
    add_body(doc, "逐项删除参考治理、扰动补偿、抗饱和或安全滤波的四个对照分别为DOB_AW_SAFE、RG_AW_SAFE、RG_DOB_SAFE、RG_DOB_AW。去掉抗饱和的RG_DOB_SAFE在爬升与反向工况均未在时窗内调定；补偿改善定高持续扰动，却没有缩短爬升捕获时间。答辩应按具体工况解释各模块作用。")

    page(doc, "十二 随机参数验证与不确定性")
    if (OUT / "monte_carlo_comparison.png").exists():
        figure(doc, "monte_carlo_comparison.png", "图5 30组配对随机参数与噪声下的性能分布", 15.0)
    mc_rows = []
    for m in ["LQI", "RG_AW", "DOB_AW_SAFE", full_name]:
        item = analysis["methods"][m]
        pitch = item["MaxPitch_deg"]
        mc_rows.append([m, f"{pitch['mean']:.2f} ± {pitch['sd']:.2f}", number(pitch["max"]),
                        f"{item['pitch_exceedance_count']}/30", number(item["SettlingTime_s"]["mean"])])
    add_table(doc, ["方法", "俯仰均值±SD °", "最差俯仰 °", "超10°组数", "平均调节 s"], mc_rows,
              [5.4, 3.2, 2.5, 2.0, 2.9])
    ci = full_mc["MaxPitch_deg"]["mean_ci95"]
    add_body(doc, f"完整方案平均俯仰峰值{number(full_mc['MaxPitch_deg']['mean'])}°，均值95%自助置信区间为[{number(ci[0])},{number(ci[1])}]°，最差{number(full_mc['MaxPitch_deg']['max'])}°。30组均进入误差带；平均调节时间{number(full_mc['SettlingTime_s']['mean'])} s，基线为{number(baseline_mc['SettlingTime_s']['mean'])} s。")
    add_body(doc, "统计单位是一次气动参数与噪声实现，不能把6001个时间点当成独立实验。使用10000次成对自助重采样计算均值及差值区间；只作描述和配对区间分析，不宣称p值显著或真实飞行器失效概率。未出现超10°的30组结果不能证明全不确定集合安全。")

    page(doc, "十三 数值检查与创新边界")
    add_heading(doc, "复现与一致性检查", 2)
    add_body(doc, f"原45 s数据独立复现，与原报告数值一致。改进方案在无噪声标称工况用0.01 s半步长复算，与0.02 s结果相比高度最大差为{number(integration['MaxAltitudeDifference_m'],6)} m，俯仰最大差为{number(integration['MaxPitchDifference_deg'],6)}°。此检查只支持该工况积分精度，不等同全部工况的收敛证明。")
    add_body(doc, "标称LQI闭环极点实部均为负；这不构成加入舵机饱和、安全滤波、切换参考及模型误差后全局稳定的证明。所有工况实际舵角与舵速均按各自限制实现，并同时记录原始命令饱和、安全滤波干预与预测不可行行为。")
    add_heading(doc, "本项目具体贡献", 2)
    add_body(doc, "第一，针对原模型的E与B不共线特点，求解高度输出调节方程生成扰动稳态补偿系数，而非套用四旋翼自抗扰控制律。第二，将俯仰对舵机速度的三阶约束转化为单输入区间投影，并将安全修正后的可实现输入反馈给积分回算。第三，在固定原LQR增益下进行九组消融与30组配对随机验证，量化姿态收益与速度代价。")
    add_body(doc, "这些贡献是具体系统上的算法设计与实验改进。扰动观测、参考治理、抗饱和和障碍函数本身均有文献基础，因此不能声称发明了这些方法，也不能仅凭组合结构就宣称国际首创。报告中的图、指标和补偿系数来自本项目代码运行。")
    add_heading(doc, "当前限制", 2)
    add_body(doc, f"完整方案在高度捕获与反向工况中参考治理出现不可行回退，最终30组中{full_mc['governor_infeasible_count']}组出现回退；安全滤波最终30组中{full_mc['safety_infeasible_count']}组出现不可行。回退计数应在答辩中解释。内部6°阈值存在微小超出，10°阈值仅在当前测试中未越界，不能写成严格硬约束保证。")
    add_body(doc, "后续优先补充真实机型气动导数、空速与迎角状态，在非线性六自由度模型中复验；给不匹配模型误差设置可证上界，设计鲁棒安全集；结合飞行任务优化6°裕量与捕获速度。物理模型尚未完善时，应将作品表述为纵向控制算法虚拟验证。")

    page(doc, "十四 结论与复现实验附件")
    add_body(doc, "本项目在原积分 LQR 高度姿态控制基础上，完成了非匹配扰动观测补偿、参考治理、执行器一致抗饱和与俯仰安全滤波的推导和实现。仿真显示，在当前纵向模型和测试参数范围内，完整方案明显压低俯仰峰值，代价是较大的目标高度误差积分和较慢的捕获过程。随机试验支持采用多层控制进行姿态风险管理，但尚不足以推断真实机型全包线能力。")
    add_table(doc, ["附件", "用途"], [
        ["simulate_lqr_aircraft_control.m", "原始45秒积分LQR基线"],
        ["simulate_improved_lqr_aircraft_control.m", "全部工况、消融、30组随机验证及半步长复算"],
        ["analyze_improved_lqr_results.py", "配对统计与10000次自助重采样"],
        ["output/improved_lqr/", "原始时间序列、参数抽样、指标、增益、图及运行配置"],
        ["build_improved_report.py", "从计算结果重建本报告，避免手填统计值"],
    ], [8.2, 7.8])
    add_heading(doc, "运行方法", 2)
    equation(doc, "simulate_improved_lqr_aircraft_control('output/improved_lqr', 30)")
    add_body(doc, "在项目目录内运行MATLAB命令，随后运行Python统计脚本与报告生成脚本。所用MATLAB版本为R2025a。原报告中的R2025b平台描述已按本次实际运行环境更正。代码仅依赖Control System Toolbox，参考治理与安全投影不依赖Optimization Toolbox。")
    add_body(doc, "正式提交前补齐队名、单位、人员与分工，确认比赛要求的PDF及附件格式。参考文献与各模块推导的对应关系见第三节，完整时间序列和指标可由附件脚本重新生成。")

    page(doc, "参考文献")
    references = [
        "[1] Stevens B L, Lewis F L, Johnson E N. Aircraft Control and Simulation 3rd ed. Wiley, 2015.",
        "[2] Etkin B, Reid L D. Dynamics of Flight Stability and Control 3rd ed. Wiley, 1996.",
        "[3] 张小明, 于纪言, 王坤坤. 自抗扰PID四旋翼飞行器控制方法研究[J]. 电子技术应用, 2019, 45(3): 84-87. DOI: 10.16157/j.issn.0258-7998.183144. 公开正文: https://www.chinaaet.com/article/3000099517",
        "[4] Chen W H, Yang J, Guo L, Li S. Disturbance-Observer-Based Control and Related Methods—An Overview[J]. IEEE Transactions on Industrial Electronics, 2016, 63(2): 1083-1095. DOI: 10.1109/TIE.2015.2478397.",
        "[5] Garone E, Di Cairano S, Kolmanovsky I. Reference and command governors for systems with constraints: A survey on theory and applications[J]. Automatica, 2017, 75: 306-328. DOI: 10.1016/j.automatica.2016.08.013.",
        "[6] Kothare M V, Campo P J, Morari M, Nett C N. A unified framework for the study of anti-windup designs[J]. Automatica, 1994, 30(12): 1869-1883. DOI: 10.1016/0005-1098(94)90048-5.",
        "[7] Xiao W, Belta C. High-Order Control Barrier Functions[J]. IEEE Transactions on Automatic Control, 2022, 67(7): 3655-3662. DOI: 10.1109/TAC.2021.3105491.",
        "[8] MathWorks. lqr Linear quadratic regulator design[EB/OL]. https://www.mathworks.com/help/control/ref/lqr.html. 访问日期: 2026-10-09.",
    ]
    for reference in references:
        p = add_body(doc, reference, first_line=False)
        p.paragraph_format.space_after = Pt(9)
        p.paragraph_format.keep_together = True

    # Prevent rows from splitting, repeat headers, and keep table captions readable.
    for table in doc.tables:
        table.autofit = False
        header_pr = table.rows[0]._tr.get_or_add_trPr()
        header_pr.append(OxmlElement("w:tblHeader"))
        for row in table.rows:
            row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
            for cell in row.cells:
                for p in cell.paragraphs:
                    p.paragraph_format.line_spacing = 1.1
                    p.paragraph_format.space_before = Pt(3)
                    p.paragraph_format.space_after = Pt(3)
                    p.paragraph_format.keep_together = True
    doc.core_properties.title = "飞行器高度姿态控制算法设计与仿真项目报告书"
    doc.core_properties.subject = "积分LQR 扰动观测 俯仰包线约束及配对仿真"
    doc.save(DEST)
    print(DEST)


if __name__ == "__main__":
    build()
