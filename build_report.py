from pathlib import Path
import csv

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
OUT = ROOT / "output"
DOCX_PATH = ROOT / "飞行器控制算法设计项目报告书初稿.docx"


def set_font(run, size=None, bold=None):
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def set_para_font(paragraph, size=11, bold=None):
    for run in paragraph.runs:
        set_font(run, size=size, bold=bold)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_borders(cell, color="D9D9D9", size="8"):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:" + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def remove_paragraph_borders(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag == qn("w:pBdr"):
            p_pr.remove(child)


def remove_style_borders(style):
    p_pr = style._element.get_or_add_pPr()
    for child in list(p_pr):
        if child.tag == qn("w:pBdr"):
            p_pr.remove(child)


def add_heading(doc, text, level=1):
    p = doc.add_heading(text, level=level)
    for run in p.runs:
        set_font(run, size=15 if level == 1 else 12, bold=True)
        run.font.color.rgb = RGBColor(0, 0, 0)
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(5)
    return p


def add_body(doc, text, first_line=True):
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.35
    p.paragraph_format.space_after = Pt(5)
    if first_line:
        p.paragraph_format.first_line_indent = Cm(0.74)
    r = p.add_run(text)
    set_font(r, size=11)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(8)
    r = p.add_run(text)
    set_font(r, size=10)
    return p


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    hdr = table.rows[0].cells
    for i, text in enumerate(headers):
        hdr[i].text = text
        set_cell_shading(hdr[i], "1F4E79")
        set_cell_borders(hdr[i])
        hdr[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for p in hdr[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_para_font(p, size=10, bold=True)
            for run in p.runs:
                run.font.color.rgb = RGBColor(255, 255, 255)
    for row in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row):
            cells[i].text = str(text)
            set_cell_borders(cells[i])
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if len(table.rows) % 2 == 1:
                set_cell_shading(cells[i], "F3F7FB")
            for p in cells[i].paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i == 0 else WD_ALIGN_PARAGRAPH.LEFT
                set_para_font(p, size=9.5)
    if widths:
        for row in table.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Cm(width)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def read_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        return list(csv.DictReader(f))[0]


def metric(value, digits=3):
    return f"{float(value):.{digits}f}"


summary = read_csv(OUT / "simulation_summary.csv")
gain = read_csv(OUT / "lqr_gain.csv")

doc = Document()
section = doc.sections[0]
section.top_margin = Cm(2.3)
section.bottom_margin = Cm(2.1)
section.left_margin = Cm(2.45)
section.right_margin = Cm(2.45)

styles = doc.styles
styles["Normal"].font.name = "SimSun"
styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
styles["Normal"]._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
styles["Normal"]._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.style = doc.styles["Title"]
remove_style_borders(doc.styles["Title"])
remove_paragraph_borders(title)
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run("飞行器控制算法设计项目报告书初稿")
set_font(r, size=20, bold=True)
r.font.color.rgb = RGBColor(0, 0, 0)

subtitle = doc.add_paragraph()
subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
subtitle.paragraph_format.space_after = Pt(18)
r = subtitle.add_run("中南大学第七届先进飞行器虚拟仿真大赛 选题五")
set_font(r, size=12)

info_rows = [
    ["项目名称", "基于积分 LQR 与指令滤波的飞行器高度姿态控制算法设计"],
    ["参赛方向", "飞行器控制算法设计"],
    ["仿真平台", "MATLAB R2025b"],
    ["当前版本", "初稿，可继续补充分工、真实型号参数与更多工况"],
]
add_table(doc, ["项目条目", "内容"], info_rows, [3.0, 12.0])

add_heading(doc, "摘要", 1)
add_body(doc, "本项目面向先进飞行器在高度捕获和姿态稳定任务中的控制需求，设计了一种基于纵向小扰动模型的积分 LQR 控制算法。算法在常规状态反馈的基础上引入高度误差积分环节，用于削弱模型偏差和持续扰动造成的稳态误差；同时加入一阶指令滤波环节，用于限制高度阶跃指令引起的舵面饱和和过大俯仰响应。")
add_body(doc, f"在 MATLAB 仿真中，飞行器执行 100 m 高度捕获任务，并在 14 s 至 20 s 之间叠加下沉阵风扰动。仿真结果表明，系统最终稳态误差约为 {metric(summary['SteadyError_m'])} m，调节时间约为 {metric(summary['SettlingTime_s'])} s，升降舵最大偏转约为 {metric(summary['MaxElevator_deg'])}°，未超过设定的 25°限幅。该结果说明所设计控制器具备基本的高度跟踪能力和扰动恢复能力，但俯仰角峰值仍需在后续工作中通过约束优化或增益调度进一步降低。")

add_heading(doc, "一 项目背景与设计目标", 1)
add_body(doc, "飞行器控制算法是连接总体设计、气动建模、任务规划和虚拟仿真的核心环节。对于固定翼或临近空间飞行器而言，高度、速度和姿态通道之间存在耦合，外界阵风、模型参数摄动以及舵面限幅都会影响控制品质。选题五要求围绕飞行控制算法开展设计，本项目选择高度姿态控制问题作为切入点，形成一个能够在 MATLAB 中复现的控制律设计与仿真验证流程。")
add_body(doc, "本初稿的目标不是完成某一真实型号的飞控定型，而是建立一套结构清楚、参数可调、结果可复现的算法样机。其核心任务包括建立简化纵向动力学模型，设计积分 LQR 状态反馈控制器，加入工程上常见的指令滤波和舵面限幅，最后通过高度捕获与阵风扰动工况评估控制性能。")

doc.add_page_break()
add_heading(doc, "二 控制对象与建模假设", 1)
add_body(doc, "为突出控制算法本身，本文采用飞行器纵向小扰动模型。状态变量选取为高度偏差、垂向速度、俯仰角和俯仰角速度，控制输入为升降舵偏角。该模型保留高度运动与俯仰运动之间的主要耦合关系，适合用于算法初步验证。")
model_rows = [
    ["h", "高度偏差", "m", "控制器主要跟踪量"],
    ["h_dot", "垂向速度", "m/s", "反映爬升或下沉趋势"],
    ["theta", "俯仰角", "rad", "姿态稳定关键变量"],
    ["q", "俯仰角速度", "rad/s", "姿态阻尼相关变量"],
    ["delta_e", "升降舵偏角", "rad", "控制输入，设置 ±25°限幅"],
]
add_table(doc, ["符号", "含义", "单位", "说明"], model_rows, [2.0, 3.2, 2.1, 7.1])
add_body(doc, "状态空间模型写为 x_dot = A x + B delta_e，其中 x = [h, h_dot, theta, q]^T。为使报告和仿真文件对应，本文采用的矩阵参数如下：")
add_body(doc, "A = [[0, 1, 0, 0], [0, -0.42, 8.50, 0.35], [0, 0, 0, 1], [0, -0.030, -1.85, -1.20]]，B = [0, 0.85, 0, 2.40]^T。")
add_body(doc, "该参数组为控制算法验证模型，主要用于反映飞行器纵向运动的耦合特性和闭环控制趋势。后续若获得具体飞行器气动导数，可直接替换 A、B 矩阵并重新整定权重。")

add_heading(doc, "三 控制算法设计", 1)
add_heading(doc, "1 积分增广状态", 2)
add_body(doc, "单纯 LQR 状态反馈能够改善动态响应，但当存在持续扰动或模型不确定性时，输出可能产生稳态误差。因此本文引入高度误差积分状态 xi，其动态方程为 xi_dot = r_h - h，其中 r_h 为经过指令滤波后的高度指令。增广状态为 x_a = [h, h_dot, theta, q, xi]^T。")
add_heading(doc, "2 LQR 性能指标", 2)
add_body(doc, "控制律通过最小化二次型性能指标获得。性能指标同时惩罚高度误差、垂向速度、俯仰角、俯仰角速度、积分误差和升降舵输入。仿真采用 Q = diag(0.90, 0.30, 45.0, 6.0, 0.16)，R = 1.15。较大的俯仰角权重用于抑制姿态偏离，输入权重用于避免舵面过度动作。")
gain_rows = [
    ["K_h", metric(gain["K_h"], 4)],
    ["K_hdot", metric(gain["K_hdot"], 4)],
    ["K_theta", metric(gain["K_theta"], 4)],
    ["K_q", metric(gain["K_q"], 4)],
    ["K_integral", metric(gain["K_integral"], 4)],
]
add_table(doc, ["反馈增益", "数值"], gain_rows, [4.0, 4.0])
add_body(doc, "最终控制律为 delta_e = sat(-K x_a)，其中 sat 表示升降舵限幅函数。仿真中升降舵限制为 ±25°。")
add_heading(doc, "3 指令滤波与约束处理", 2)
add_body(doc, "若直接给定 100 m 阶跃高度指令，控制器容易在初始阶段给出过大的姿态和舵面动作。本文在参考输入前加入一阶指令滤波 r_h(t) = 100(1 - exp(-t/7))，相当于对飞行器提出更平滑的高度捕获任务。该处理属于参考治理方法，能够在不改变主反馈结构的前提下降低饱和风险。")

add_heading(doc, "四 MATLAB 仿真方案", 1)
sim_rows = [
    ["仿真时间", "0 到 45 s"],
    ["高度指令", "100 m，经 7 s 一阶滤波"],
    ["扰动设置", "14 s 到 20 s 加入 -0.80 m/s^2 垂向阵风加速度"],
    ["舵面限制", "升降舵 ±25°"],
    ["求解器", "ode45，RelTol = 1e-7，AbsTol = 1e-9"],
    ["输出文件", "高度响应图、姿态与舵面图、控制增益 CSV、性能指标 CSV"],
]
add_table(doc, ["项目", "设置"], sim_rows, [4.0, 11.0])
add_body(doc, "仿真脚本保存为 simulate_lqr_aircraft_control.m。脚本首先建立增广状态空间模型并调用 lqr 求解反馈增益，然后通过 ode45 对含饱和与阵风扰动的闭环系统进行数值积分，最后导出 PNG 图像和 CSV 数据。")

add_heading(doc, "五 仿真结果与分析", 1)
metric_rows = [
    ["高度指令", f"{metric(summary['RefAltitude_m'], 1)} m"],
    ["最大超调", f"{metric(summary['Overshoot_m'])} m"],
    ["调节时间", f"{metric(summary['SettlingTime_s'])} s"],
    ["稳态误差", f"{metric(summary['SteadyError_m'])} m"],
    ["最大俯仰角", f"{metric(summary['MaxPitch_deg'])}°"],
    ["最大升降舵偏角", f"{metric(summary['MaxElevator_deg'])}°"],
    ["绝对误差积分", f"{metric(summary['IAE_m_s'])} m·s"],
]
add_table(doc, ["性能指标", "仿真结果"], metric_rows, [5.0, 6.0])
add_body(doc, "从高度响应可以看出，指令滤波后闭环系统能够平滑接近 100 m 高度目标，阵风扰动期间高度曲线出现轻微偏离，但扰动结束后能够继续收敛。最终稳态误差小于 0.4 m，说明积分环节有效提高了高度跟踪精度。")

if (OUT / "altitude_response.png").exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OUT / "altitude_response.png"), width=Cm(14.0))
    add_caption(doc, "图 1 高度跟踪响应与阵风扰动窗口")

add_body(doc, "姿态与舵面响应表明，升降舵最大偏角约为 24.19°，低于 25°限幅，说明参考治理后控制输入没有持续饱和。最大俯仰角约为 29.10°，对于常规小角度巡航控制仍偏大，说明后续应进一步提高俯仰角权重、引入俯仰角硬约束，或采用模型预测控制等约束优化算法。")

if (OUT / "control_states.png").exists():
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(OUT / "control_states.png"), width=Cm(14.0))
    add_caption(doc, "图 2 俯仰状态与升降舵控制量")

add_heading(doc, "六 方案特点与创新点", 1)
add_body(doc, "第一，控制器采用积分增广结构，将高度跟踪精度和抗持续扰动能力纳入同一状态反馈框架。第二，算法加入指令滤波，能够降低阶跃命令造成的舵面饱和风险，使仿真更接近实际飞控系统中的指令管理逻辑。第三，仿真脚本自动输出图像、增益和性能指标，便于后续开展权重扫描、鲁棒性分析和报告复现。")
add_body(doc, "与仅展示理论控制律相比，本方案更强调从模型、控制器、限幅、扰动到结果分析的完整闭环流程。后续可以继续扩展为多工况批量仿真、参数不确定性蒙特卡洛分析，以及与 PID、极点配置或模型预测控制方法的对比。")

doc.add_page_break()
add_heading(doc, "七 后续优化计划", 1)
plan_rows = [
    ["模型完善", "补充真实飞行器质量、气动导数、舵机一阶动态和传感器噪声模型"],
    ["算法优化", "开展 Q、R 权重扫描，降低俯仰角峰值和绝对误差积分"],
    ["对比验证", "增加 PID、极点配置和 MPC 对比组，形成算法优劣分析"],
    ["鲁棒性测试", "加入不同阵风强度、初始偏差、参数摄动和舵面迟滞工况"],
    ["作品附件", "整理 MATLAB 脚本、仿真图、演示视频和最终版研究报告"],
]
add_table(doc, ["工作方向", "具体内容"], plan_rows, [4.0, 11.0])

add_heading(doc, "八 结论", 1)
add_body(doc, "本文围绕先进飞行器虚拟仿真大赛选题五，完成了飞行器高度姿态控制算法的初步方案设计。基于简化纵向小扰动模型，本文构建了积分 LQR 控制律，并通过指令滤波和舵面限幅提高了闭环系统的工程可用性。MATLAB 仿真结果显示，系统能够在阵风扰动下完成 100 m 高度捕获任务，最终稳态误差较小，控制输入满足限幅要求。")
add_body(doc, "当前方案仍属于初稿，主要价值在于形成了可复现的算法框架和仿真基线。下一阶段应结合更真实的飞行器参数和更多工况开展整定，使控制器在姿态约束、响应速度和鲁棒性之间达到更优平衡。")

add_heading(doc, "参考文献", 1)
refs = [
    "1. Stevens B L, Lewis F L, Johnson E N. Aircraft Control and Simulation. Wiley, 2015.",
    "2. Etkin B, Reid L D. Dynamics of Flight Stability and Control. Wiley, 1996.",
    "3. Bryson A E, Ho Y C. Applied Optimal Control. Taylor and Francis, 1975.",
    "4. MathWorks. Linear Quadratic Regulator LQR Design Documentation.",
]
for item in refs:
    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.25
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(item)
    set_font(r, size=10.5)

doc.save(DOCX_PATH)
print(DOCX_PATH)
